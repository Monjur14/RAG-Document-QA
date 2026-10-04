"""Builds small PDF/DOCX fixtures in memory for parser tests."""
import io
import zipfile


def make_pdf(blank_page: bool = False) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4)
    st = getSampleStyleSheet()
    if blank_page:
        doc.build([Spacer(1, 10)])
        return buf.getvalue()
    table = Table([["Item", "Value"], ["RAM", "2 GB"], ["Disk", "10 GB"]])
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.8, colors.black)]))
    body = "Orbit runs scheduled jobs on a worker pool and retries failures automatically. " * 3
    story = [
        Paragraph("Installation Guide", st["Heading1"]),
        Paragraph(body, st["BodyText"]),
        Spacer(1, 14),
        Paragraph("Requirements", st["Heading2"]),
        Paragraph("The minimum hardware requirements are listed below. " + body, st["BodyText"]),
        Spacer(1, 10),
        table,
        PageBreak(),
        Paragraph("Troubleshooting continues on the second page of this guide. " + body, st["BodyText"]),
        Paragraph("Support", st["Heading2"]),
        Paragraph("Contact the support team through the community forum. " + body, st["BodyText"]),
    ]
    doc.build(story)
    return buf.getvalue()


def make_docx() -> bytes:
    from docx import Document

    d = Document()
    d.add_heading("Orbit Handbook", level=1)
    d.add_paragraph("Orbit schedules jobs.")
    d.add_heading("Limits", level=2)
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Plan"
    t.cell(0, 1).text = "Jobs"
    t.cell(1, 0).text = "Free"
    t.cell(1, 1).text = "100"
    d.add_paragraph("")
    d.add_paragraph("Limits reset monthly.")
    merged = d.add_table(rows=1, cols=3)
    merged.cell(0, 0).merge(merged.cell(0, 1)).text = "Merged"
    merged.cell(0, 2).text = "Solo"
    d.sections[0].header.paragraphs[0].text = "HEADER TEXT"
    d.sections[0].footer.paragraphs[0].text = "FOOTER TEXT"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_zip_without_document_xml() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("hello.txt", "hi")
    return buf.getvalue()


def make_zip_bomb_docx(megabytes: int = 3) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", "<x/>")
        z.writestr("word/big.bin", b"\0" * megabytes * 1024 * 1024)
    return buf.getvalue()
