"""Builds the two documents that are written for this project (so they are free to publish with the repo):

    python -m evals.make_own_corpus

  harbor-handbook.docx   a Word guide with heading styles, a bulleted list, a numbered list and two tables
  harbor-faq.html        a web page with navigation/footer boilerplate, a table, a list and one hidden element

"Harbor" is an invented backup service; nothing here is real. They exist so the per-format evals have a real Word
file and a real HTML file with labelled questions (evals/data/corpus_questions_3.json). The HTML page also holds a
hidden decoy fact that the parser must drop: the unanswerable question u21 asks about it.
"""
from pathlib import Path

from evals.corpus_eval import CORPUS

HTML = """<!doctype html>
<html lang="en"><head><title>Harbor FAQ</title><style>body{font-family:sans-serif}</style></head>
<body>
<nav><a href="/">Home</a> <a href="/pricing">Pricing</a> <a href="/faq">FAQ</a></nav>
<header>Harbor Backup Help Centre</header>
<main>
<h1>Frequently asked questions</h1>
<h2>Backups</h2>
<p>External drives are backed up only while they are connected during the scheduled run.</p>
<p>Temporary files ending in .tmp and the system swap file are skipped by default.</p>
<p>Upload speed is limited to 50 percent of the available bandwidth during working hours.</p>
<h2>Accounts and billing</h2>
<p>One account can protect up to 5 devices on the Starter plan and up to 25 devices on the Pro plan.</p>
<p>To cancel a subscription, open Billing and press Cancel plan; access to existing backups continues for 14 days.</p>
<p>Deleted files stay in the history for 30 days on the Starter plan.</p>
<h2>Data centres</h2>
<table>
<tr><th>Region</th><th>Data centre</th><th>Latency target</th></tr>
<tr><td>Europe</td><td>Frankfurt</td><td>under 40 ms</td></tr>
<tr><td>Asia</td><td>Singapore</td><td>under 60 ms</td></tr>
</table>
<ul>
<li>Every data centre is certified for ISO 27001.</li>
<li>Customers cannot choose a data centre after the first backup has been uploaded.</li>
</ul>
<div style="display:none">Harbor offers lifetime free storage of 10 TB to every new customer.</div>
</main>
<footer>Copyright Harbor. All rights reserved.</footer>
</body></html>
"""


def build_docx() -> bytes:
    import io

    from docx import Document

    d = Document()
    d.add_heading("Harbor Backup Administrator Guide", level=0)
    d.add_heading("Overview", level=1)
    d.add_paragraph("Harbor is a backup service that copies folders to encrypted storage every night at 01:30 server time. "
                    "Each backup is split into blocks of 4 MB before upload.")
    d.add_heading("Plans", level=1)
    plans = [("Plan", "Storage", "Retention", "Price"), ("Starter", "500 GB", "30 days", "8 USD per month"),
             ("Pro", "5 TB", "90 days", "25 USD per month"), ("Team", "20 TB", "365 days", "90 USD per month")]
    t = d.add_table(rows=len(plans), cols=4)
    for r, row in enumerate(plans):
        for c, val in enumerate(row):
            t.cell(r, c).text = val
    d.add_paragraph("All plans include end-to-end encryption and a restore test once a month.")
    d.add_heading("Installation", level=1)
    d.add_heading("Requirements", level=2)
    for item in ["Windows 10 or newer, macOS 12 or newer, or Ubuntu 22.04.",
                 "At least 2 GB of free memory.",
                 "A network connection that allows outbound traffic on port 8443."]:
        d.add_paragraph(item, style="List Bullet")
    d.add_heading("Steps", level=2)
    for item in ["Download the installer from the account page.",
                 "Run the installer and sign in with your account email.",
                 "Pick the folders to protect and press Start."]:
        d.add_paragraph(item, style="List Number")
    d.add_heading("Restoring files", level=1)
    d.add_paragraph("To restore a deleted file, open the History tab, choose a date, and press Restore next to the file name.")
    d.add_paragraph("Restored files are placed in a folder named Harbor-Restored on the desktop unless another location is chosen.")
    d.add_paragraph("Restoring a full machine takes about one hour for every 100 GB of data.")
    d.add_heading("Security", level=1)
    d.add_paragraph("Encryption keys never leave the device. If the key is lost, the data cannot be recovered, even by support staff.")
    d.add_paragraph("Administrators can enable two-factor authentication under Settings, then Security.")
    d.add_heading("Troubleshooting", level=1)
    errs = [("Error code", "Meaning", "Fix"),
            ("E101", "Network unreachable", "Check that the firewall allows port 8443."),
            ("E204", "Disk full on target", "Free at least 10 percent of the disk."),
            ("E307", "Key mismatch", "Re-import the recovery key from the printed sheet.")]
    t = d.add_table(rows=len(errs), cols=3)
    for r, row in enumerate(errs):
        for c, val in enumerate(row):
            t.cell(r, c).text = val
    d.add_heading("Support", level=1)
    d.add_paragraph("Support is available on weekdays from 08:00 to 18:00 UTC through the in-app chat.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def main() -> int:
    CORPUS.mkdir(parents=True, exist_ok=True)
    (CORPUS / "harbor-handbook.docx").write_bytes(build_docx())
    (CORPUS / "harbor-faq.html").write_text(HTML, encoding="utf-8")
    print(f"wrote harbor-handbook.docx and harbor-faq.html to {CORPUS}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
